"""Investment-committee workbook (Excel) and buy order (CSV) from a run result."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CR = 1e7
HEAD = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _sheet(writer, name, df: pd.DataFrame, widths=None):
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    for c in ws[1]:
        c.fill, c.font = HEAD, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col) or min(max(len(str(col)), *(len(f"{v}") for v in df[col].head(50)), 8) + 2, 48)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def buy_order_csv(result: dict) -> str:
    rows = [{
        "ISIN": h["isin"], "Security": h["description"], "Maturity": h["maturity"],
        "Face value (Rs)": round(h["face"], 2), "Lots": h.get("lots"),
        "Clean price": h["clean"], "Accrued": round(h["accrued"], 4), "Dirty price": round(h["dirty"], 4),
        "Estimated cost (Rs)": round(h["cost"], 2), "Settlement": result["settle"],
    } for h in result.get("holdings", [])]
    return pd.DataFrame(rows).to_csv(index=False)


def workbook(result: dict) -> bytes:
    s = result.get("summary", {})
    ds = result["dataset"]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        summary = [
            ("Result", "All payments covered" if s.get("covered_all") else (result.get("message") or result["status"])),
            ("Total investment (Rs cr)", s.get("cost", 0) / CR),
            ("Liabilities to pay (Rs cr)", s.get("required_total", 0) / CR),
            ("PV of liabilities on the G-Sec curve (Rs cr)", s.get("liability_pv", 0) / CR),
            ("Saving vs manual selection (Rs lakh)", (s.get("savings_vs_manual") or 0) / 1e5),
            ("Bonds used", s.get("bonds_used")),
            ("Lowest coverage of any payment", s.get("min_coverage")),
            ("Cash left after final payment (Rs)", s.get("leftover_cash")),
            ("Portfolio modified duration (yrs)", s.get("portfolio_mod_dur")),
            ("Liability modified duration (yrs)", s.get("liability_mod_dur")),
            ("Portfolio convexity", s.get("portfolio_convexity")),
            ("Portfolio yield (%)", s.get("portfolio_yield")),
            ("Average days cash waits before use", s.get("avg_wait_days")),
            ("Data source", ds["label"]),
            ("Last trade date", ds["last_trade_date"]),
            ("Settlement date (T+1)", result["settle"]),
        ]
        _sheet(w, "Summary", pd.DataFrame(summary, columns=["Item", "Value"]), {"Item": 46, "Value": 48})

        hold = pd.DataFrame(result.get("holdings", []))
        if len(hold):
            hold = hold[["isin", "description", "category", "maturity", "coupon", "face", "lots", "clean",
                         "accrued", "dirty", "cost", "weight", "yield", "mod_dur", "volume", "trades",
                         "order_pct_volume", "liquidity"]]
            hold.columns = ["ISIN", "Security", "Type", "Maturity", "Coupon %", "Face (Rs)", "Lots", "Clean",
                            "Accrued", "Dirty", "Cost (Rs)", "Weight", "Yield %", "Mod. duration",
                            "Day volume (Rs cr)", "Trades", "Order % of volume", "Liquidity"]
            _sheet(w, "Buy list", hold)

        sch = pd.DataFrame(result.get("schedule", []))
        if len(sch):
            sch = sch[["date", "liability", "required", "existing_cash", "bond_cash", "carried_in", "paid",
                       "shortfall", "carried_out", "coverage"]]
            sch.columns = ["Payment date", "Liability", "Required (with margin)", "Existing holdings cash",
                           "New bond cash", "Carried in", "Paid", "Shortfall", "Carried out", "Coverage"]
            _sheet(w, "Coverage", sch)

        if result.get("duals"):
            d = pd.DataFrame(result["duals"])[["date", "years", "shadow_price", "implied_rate"]]
            d.columns = ["Payment date", "Years from settlement", "Cost today of Rs 1 due", "Implied rate %"]
            _sheet(w, "Shadow prices", d)

        if result.get("shock"):
            sh = pd.DataFrame(result["shock"])
            sh.columns = ["Shift (bp)", "Portfolio value", "Portfolio change %", "Duration estimate %",
                          "Duration+convexity estimate %", "Liability PV", "Liability change %", "Surplus"]
            _sheet(w, "Rate risk", sh)

        if result.get("scenarios"):
            sc = pd.DataFrame([{k: v for k, v in r.items() if k != "units"} for r in result["scenarios"]])
            _sheet(w, "Scenarios", sc)

        if result.get("strategies"):
            st = pd.DataFrame([{k: v for k, v in r.items() if k != "holdings"} for r in result["strategies"]])
            _sheet(w, "Strategies", st)

        uni = pd.DataFrame(result.get("universe", []))
        if len(uni):
            _sheet(w, "Eligible bonds", uni)
        exc = pd.DataFrame(result.get("excluded", []))
        if len(exc):
            _sheet(w, "Excluded", exc)

        assumptions = [
            "Bond data: CCIL Trading Platform, weighted-average clean price of the latest trading day per ISIN ('All' platform).",
            "Coupon parsed from the security description; floating-rate, gold, inflation-indexed bonds and STRIPS excluded.",
            "G-Secs and SDLs: semi-annual coupons on the maturity day/month, 30/360, face value Rs 100. T-Bills: zero coupon.",
            "Cost = dirty price (clean + accrued to T+1 settlement) plus transaction cost.",
            "Network LP: one node per payment date; bond cash flows into the next payment date; surplus carried forward "
            f"earning {result['options'].get('reinvest_rate', 0)}% a year.",
            "Liability PV and duration use a Nelson-Siegel zero curve fitted to sovereign bond prices.",
            "Rate scenarios shift every bond's yield in parallel and re-price from its cash flows.",
            "CCIL does not authorise commercial use of its website data without written permission.",
        ]
        _sheet(w, "Assumptions", pd.DataFrame({"Assumption": assumptions}), {"Assumption": 120})
    return buf.getvalue()
